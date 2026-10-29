"""Excel export of the Tracking page (exactly the rows the user is looking at)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="E7E9F0")
_BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)

# status text -> (fill, font colour)
_STATUS_STYLE = {
    "Run in last 7 days":   ("DCFCE7", "166534"),
    "Run 8-30 days ago":    ("FEF3C7", "92400E"),
    "Not run in 30+ days":  ("FEE2E2", "991B1B"),
    "Never run":            ("EEF0F6", "5B6178"),
}


def _text(value) -> str:
    return "" if value is None else str(value)


def build_tracking_xlsx(rows: list[dict], summary: dict) -> bytes:
    wb = Workbook()

    # ---- sheet 1: one row per screen ------------------------------------------------
    ws = wb.active
    ws.title = "Tracking"
    headers = ["Module", "Screen", "Last run", "Days since last run", "Runs (last 30 days)", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    for r in rows:
        days = r.get("days")
        ws.append([
            _text(r.get("module")), _text(r.get("screen")), _text(r.get("last_run_text")) or "Never run",
            days if isinstance(days, (int, float)) else "",
            r.get("runs30") if isinstance(r.get("runs30"), (int, float)) else 0,
            _text(r.get("status")),
        ])
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center")
        for col in (4, 5):
            ws.cell(row=row_idx, column=col).alignment = Alignment(horizontal="center", vertical="center")
        style = _STATUS_STYLE.get(_text(r.get("status")))
        if style:
            status_cell = ws.cell(row=row_idx, column=6)
            status_cell.fill = PatternFill("solid", fgColor=style[0])
            status_cell.font = Font(color=style[1], bold=True)

    for idx, width in enumerate([28, 38, 24, 20, 20, 22], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # ---- sheet 2: summary -----------------------------------------------------------
    sm = wb.create_sheet("Summary")
    lines = [
        ("QC run tracking", ""),
        ("Created", _text(summary.get("generated"))),
        ("Showing", _text(summary.get("showing")) or "All screens"),
        ("Search", _text(summary.get("search")) or "-"),
        ("", ""),
        ("Screens in this export", len(rows)),
        ("Run in last 7 days", summary.get("fresh", 0)),
        ("Not run in 30+ days", summary.get("stale", 0)),
        ("Never run", summary.get("never", 0)),
        ("Coverage (screens run in last 30 days)", _text(summary.get("coverage"))),
    ]
    for label, value in lines:
        sm.append([label, value])
    sm["A1"].font = Font(bold=True, size=14)
    for row in sm.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)
    sm.column_dimensions["A"].width = 42
    sm.column_dimensions["B"].width = 36

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()